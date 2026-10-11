"""
将 user_feedback_report_124.md 转换为 Word 文档
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re

# 读取 Markdown 文件
md_path = r"c:\Users\11127\Desktop\李泓毅新版\AI-mental-main\algorithm\artifacts\metrics\user_feedback_report_124.md"
with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)

# 页面边距
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

def add_heading_styled(doc, text, level):
    """添加带样式的标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = '微软雅黑'
    return heading

def add_table_from_lines(doc, lines, start_idx):
    """从 Markdown 表格行解析并添加表格"""
    # 找到表格的结束位置
    table_lines = []
    idx = start_idx
    while idx < len(lines) and lines[idx].strip().startswith('|'):
        table_lines.append(lines[idx].strip())
        idx += 1
    
    if len(table_lines) < 2:
        return idx
    
    # 解析表头
    headers = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
    # 跳过分隔行
    data_start = 2 if len(table_lines) > 1 and '---' in table_lines[1] else 1
    
    # 创建表格
    num_rows = len(table_lines) - data_start
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows + 1, cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 填充表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
    
    # 填充数据
    for row_idx, line in enumerate(table_lines[data_start:]):
        cells = [cell.strip() for cell in line.split('|')[1:-1]]
        for col_idx, cell_text in enumerate(cells):
            if col_idx < num_cols:
                cell = table.rows[row_idx + 1].cells[col_idx]
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    
    doc.add_paragraph()  # 表格后空行
    return idx

# 逐行解析
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    
    # 一级标题
    if line.startswith('# ') and not line.startswith('## '):
        title_text = line[2:].strip()
        add_heading_styled(doc, title_text, level=0)
        i += 1
        continue
    
    # 二级标题
    if line.startswith('## '):
        add_heading_styled(doc, line[3:].strip(), level=1)
        i += 1
        continue
    
    # 三级标题
    if line.startswith('### '):
        add_heading_styled(doc, line[4:].strip(), level=2)
        i += 1
        continue
    
    # 四级标题
    if line.startswith('#### '):
        add_heading_styled(doc, line[5:].strip(), level=3)
        i += 1
        continue
    
    # 表格
    if line.strip().startswith('|'):
        i = add_table_from_lines(doc, lines, i)
        continue
    
    # 无序列表
    if line.strip().startswith('- ') or line.strip().startswith('* '):
        text = line.strip()[2:]
        # 处理加粗
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        p = doc.add_paragraph(text, style='List Bullet')
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1
        continue
    
    # 有序列表
    m = re.match(r'^(\d+)\.\s+(.+)', line.strip())
    if m:
        text = m.group(2)
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        p = doc.add_paragraph(text, style='List Number')
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1
        continue
    
    # 引用块
    if line.strip().startswith('> '):
        text = line.strip()[2:]
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        p = doc.add_paragraph(text)
        p.paragraph_format.left_indent = Inches(0.5)
        for run in p.runs:
            run.font.size = Pt(10)
            run.italic = True
        i += 1
        continue
    
    # 分隔线
    if line.strip() == '---' or line.strip() == '***':
        doc.add_paragraph()
        i += 1
        continue
    
    # 普通段落
    if line.strip():
        text = line.strip()
        # 处理加粗标记
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        p = doc.add_paragraph(text)
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1
        continue
    
    # 空行
    i += 1

# 保存
output_path = r"c:\Users\11127\Desktop\李泓毅新版\AI-mental-main\algorithm\artifacts\metrics\user_feedback_report_124.docx"
doc.save(output_path)
print(f"已生成 Word 文档：{output_path}")
